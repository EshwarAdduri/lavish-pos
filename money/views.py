import csv
import io
from datetime import date, timedelta
from decimal import Decimal

from django import forms
from django.contrib import messages
from django.core.paginator import Paginator
from django.db import transaction
from django.db.models import Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from core.permissions import can_record, owner_required, recorder_required
from core.utils import money, today
from orders.models import Order, OrderItem, OrderStatus, OrderType, Payment, PaymentMethod
from stock import services as stock_services
from stock.models import RecipeLine, StockCount, StockItem

from . import services
from .models import DayClose, Expense

WEEKDAYS = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]


def _date(value, default):
    try:
        return date.fromisoformat(value) if value else default
    except ValueError:
        return default


def resolve_range(request):
    t = today()
    preset = request.GET.get("range", "7d")
    if preset == "today":
        return t, t, preset
    if preset == "yesterday":
        y = t - timedelta(days=1)
        return y, y, preset
    if preset == "week":
        return t - timedelta(days=t.weekday()), t, preset
    if preset == "month":
        return t.replace(day=1), t, preset
    if preset == "lastmonth":
        end = t.replace(day=1) - timedelta(days=1)
        return end.replace(day=1), end, preset
    if preset == "30d":
        return t - timedelta(days=29), t, preset
    if preset == "custom":
        start = _date(request.GET.get("start"), t - timedelta(days=6))
        end = _date(request.GET.get("end"), t)
        if start > end:
            start, end = end, start
        if (end - start).days > 400:
            start = end - timedelta(days=400)
        return start, end, preset
    return t - timedelta(days=6), t, "7d"


def _chart(p):
    return {
        "daily": {
            "labels": [d.strftime("%d %b") for d in p["daily"]],
            "sales": [float(v["sales"]) for v in p["daily"].values()],
            "orders": [v["orders"] for v in p["daily"].values()],
        },
        "hour": {"labels": [f"{h:02d}" for h in range(24)], "sales": [float(v) for v in p["by_hour"].values()]},
        "weekday": {
            "labels": WEEKDAYS,
            "sales": [float(p["by_weekday"][d]["sales"]) for d in range(1, 8)],
            "orders": [p["by_weekday"][d]["orders"] for d in range(1, 8)],
        },
        "methods": {
            "labels": [PaymentMethod(m).label for m in p["by_method"]],
            "values": [float(v) for v in p["by_method"].values()],
        },
        "types": {
            "labels": [OrderType(r["order_type"]).label for r in p["by_type"]],
            "values": [float(r["s"] or 0) for r in p["by_type"]],
        },
    }


# ------------------------------------------------------------------ dashboard
@owner_required
def dashboard(request):
    t = today()
    s = services.day_summary(t)
    y = services.day_summary(t - timedelta(days=1))
    p = services.period(t - timedelta(days=13), t)
    stock_rows = [r for r in stock_services.day_stock(t) if r.item.is_key or r.is_low]
    stock_started = any(r.tracked for r in stock_rows)
    yesterday_closed = DayClose.objects.filter(business_date=t - timedelta(days=1)).exists()
    had_sales_yesterday = y.orders_count > 0
    return render(request, "money/dashboard.html", {
        "s": s, "y": y, "p": p, "chart": _chart(p), "stock_rows": stock_rows, "stock_started": stock_started,
        "close_reminder": had_sales_yesterday and not yesterday_closed,
        "unverified": RecipeLine.objects.filter(verified=False).count(),
        "low": [r for r in stock_rows if r.is_low],
        "recent_payments": Payment.objects.filter(business_date=t).select_related("order", "received_by").order_by("-created_at")[:8],
    })


@owner_required
def reports(request):
    start, end, preset = resolve_range(request)
    p = services.period(start, end)
    return render(request, "money/reports.html", {
        "p": p, "chart": _chart(p), "preset": preset, "start": start, "end": end,
        "exp_labels": dict(Expense.Category.choices), "type_labels": dict(OrderType.choices),
    })


# ------------------------------------------------------------------ expenses
class ExpenseForm(forms.ModelForm):
    class Meta:
        model = Expense
        fields = ["business_date", "category", "description", "amount", "paid_from"]
        widgets = {
            "business_date": forms.DateInput(attrs={"type": "date"}),
            "amount": forms.NumberInput(attrs={"step": "0.01", "inputmode": "decimal"}),
        }
        labels = {"business_date": "Date"}


def expenses(request):
    if not can_record(request.user):
        return redirect("pos")
    start, end, preset = resolve_range(request)
    qs = Expense.objects.filter(business_date__range=(start, end)).select_related("created_by")
    cat = request.GET.get("category", "")
    if cat:
        qs = qs.filter(category=cat)
    page = Paginator(qs, 50).get_page(request.GET.get("page"))
    return render(request, "money/expenses.html", {
        "page": page, "start": start, "end": end, "preset": preset, "category": cat,
        "categories": Expense.Category.choices, "total": qs.aggregate(s=Sum("amount"))["s"] or 0,
    })


@recorder_required
def expense_edit(request, pk=None):
    obj = get_object_or_404(Expense, pk=pk) if pk else None
    if obj and not request.user.is_owner:
        messages.error(request, "Only the owner can change saved expenses.")
        return redirect("expenses")
    if obj and obj.stock_movement_id:
        messages.info(request, "This expense came from a stock purchase. Delete that stock entry instead.")
        return redirect("stock_history")
    form = ExpenseForm(request.POST or None, instance=obj, initial={"business_date": today()})
    if request.method == "POST":
        if "delete" in request.POST and obj:
            obj.delete()
            messages.success(request, "Expense deleted (kept in the audit log).")
            return redirect("expenses")
        if form.is_valid():
            e = form.save(commit=False)
            if not obj:
                e.created_by = request.user
            e.save()
            messages.success(request, "Expense saved.")
            return redirect("expenses")
    return render(request, "menu/simple_form.html", {
        "form": form, "obj": obj, "title": "Expense", "back": "/money/expenses/", "can_delete": bool(obj),
    })


# ------------------------------------------------------------------ day close
@recorder_required
def day_close(request):
    d = _date(request.GET.get("date") or request.POST.get("date"), today())
    existing = DayClose.objects.filter(business_date=d).first()
    if existing and not request.user.is_owner and request.method == "POST":
        messages.error(request, "This day is already closed. Only the owner can change it.")
        return redirect(f"/money/close/?date={d}")
    summary = services.day_summary(d)
    stock_rows = [r for r in stock_services.day_stock(d) if r.item.count_daily]
    errors = []

    if request.method == "POST":
        def dec(name, default="0"):
            raw = (request.POST.get(name) or default).strip()
            try:
                v = money(Decimal(raw))
            except Exception:
                errors.append(f"“{raw}” is not a valid amount.")
                return Decimal("0")
            if v < 0:
                errors.append("Amounts cannot be negative.")
            return v

        opening = dec("opening_cash", str(summary.opening_cash))
        counted = dec("counted_cash")
        removed = dec("cash_removed")
        if removed > counted:
            errors.append("Cash taken out cannot be more than the cash counted.")
        counts = {}
        for r in stock_rows:
            raw = request.POST.get(f"count-{r.item.id}", "").strip()
            if raw == "":
                continue
            try:
                q = Decimal(raw)
                if q < 0:
                    raise ValueError
                counts[r.item.id] = q.quantize(Decimal("0.001"))
            except Exception:
                errors.append(f"Count for {r.item.name} is not a valid number.")
        if not errors:
            with transaction.atomic():
                # Recompute with the opening cash the user confirmed.
                expected = opening + summary.by_method[PaymentMethod.CASH] - summary.expenses_from_drawer
                loss = Decimal("0")
                for r in stock_rows:
                    if r.item.id not in counts:
                        continue
                    q = counts[r.item.id]
                    var = q - r.expected
                    val = money(var * r.item.cost_per_unit)
                    StockCount.objects.update_or_create(
                        business_date=d, stock_item=r.item,
                        defaults={"counted_qty": q, "expected_qty": r.expected, "variance": var,
                                  "variance_value": val, "counted_by": request.user},
                    )
                    if val < 0:
                        loss += val
                DayClose.objects.update_or_create(
                    business_date=d,
                    defaults={
                        "opening_cash": opening, "cash_sales": summary.by_method[PaymentMethod.CASH],
                        "cash_expenses": summary.expenses_from_drawer, "expected_cash": expected,
                        "counted_cash": counted, "cash_difference": counted - expected, "cash_removed": removed,
                        "upi_sales": summary.by_method[PaymentMethod.UPI], "card_sales": summary.by_method[PaymentMethod.CARD],
                        "online_sales": summary.by_method[PaymentMethod.ONLINE], "total_sales": summary.sales,
                        "unpaid_amount": summary.unpaid, "orders_count": summary.orders_count,
                        "cancelled_count": summary.cancelled_count, "total_expenses": summary.expenses,
                        "stock_loss_value": -loss, "notes": request.POST.get("notes", "")[:300],
                        "closed_by": request.user,
                    },
                )
            messages.success(request, f"Day {d:%d %b} closed.")
            return redirect(f"/money/closings/{d.isoformat()}/" if request.user.is_owner else "/pos/")

    src = request.POST if request.method == "POST" else {}
    init = {
        "opening": src.get("opening_cash") or (existing.opening_cash if existing else summary.opening_cash),
        "counted": src.get("counted_cash") or (existing.counted_cash if existing else ""),
        "removed": src.get("cash_removed") or (existing.cash_removed if existing else 0),
        "notes": src.get("notes") or (existing.notes if existing else ""),
    }
    return render(request, "money/close.html", {
        "day": d, "s": summary, "stock_rows": stock_rows, "existing": existing, "errors": errors, "init": init,
        "prev_day": d - timedelta(days=1), "next_day": d + timedelta(days=1), "is_today": d == today(),
    })


@owner_required
def closings(request):
    page = Paginator(DayClose.objects.select_related("closed_by"), 31).get_page(request.GET.get("page"))
    return render(request, "money/closings.html", {"page": page})


@owner_required
def closing_report(request, day):
    d = _date(day, today())
    close = get_object_or_404(DayClose, business_date=d)
    counts = StockCount.objects.filter(business_date=d).select_related("stock_item")
    p = services.period(d, d)
    return render(request, "money/closing_report.html", {"c": close, "counts": counts, "p": p,
                                                          "type_labels": dict(OrderType.choices)})


# ------------------------------------------------------------------ exports
def _csv_response(name, header, rows):
    buf = io.StringIO()
    buf.write("﻿")  # so Excel opens ₹ and Telugu names correctly
    w = csv.writer(buf)
    w.writerow(header)
    w.writerows(rows)
    resp = HttpResponse(buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{name}"'
    return resp


def _orders_rows(start, end):
    qs = Order.objects.filter(business_date__range=(start, end)).select_related("created_by", "device").order_by("created_at")
    for o in qs:
        yield [
            o.bill_number, o.token_no, o.business_date.isoformat(), timezone.localtime(o.created_at).strftime("%Y-%m-%d %H:%M:%S"),
            o.get_order_type_display(), o.get_platform_display() if o.platform else "", o.get_status_display(),
            float(o.subtotal), float(o.discount_amount), float(o.total), float(o.paid_amount),
            o.created_by.username, str(o.device) if o.device else "", o.customer_name, o.cancel_reason,
        ]


ORDER_HEADER = ["Bill", "Token", "Business date", "Created at", "Type", "Platform", "Status", "Subtotal",
                "Discount", "Total", "Paid", "Staff", "Device", "Customer", "Cancel reason"]
ITEM_HEADER = ["Bill", "Business date", "Hour", "Weekday", "Category", "Item", "Size", "Qty", "Unit price",
               "Extras price", "Line total", "Order type", "Status"]


def _item_rows(start, end):
    qs = OrderItem.objects.filter(order__business_date__range=(start, end)).select_related("order").order_by("order__created_at", "id")
    for i in qs:
        o = i.order
        local = timezone.localtime(o.created_at)
        yield [o.bill_number, o.business_date.isoformat(), local.hour, WEEKDAYS[o.business_date.weekday()],
               i.category_name, i.item_name, i.variant_name, i.quantity, float(i.unit_price), float(i.addons_price),
               float(i.line_total), o.get_order_type_display(), o.get_status_display()]


@owner_required
def export(request, kind):
    start, end, _ = resolve_range(request)
    stamp = f"{start}_{end}"
    if kind == "orders.csv":
        return _csv_response(f"orders_{stamp}.csv", ORDER_HEADER, _orders_rows(start, end))
    if kind == "items.csv":
        return _csv_response(f"items_{stamp}.csv", ITEM_HEADER, _item_rows(start, end))
    if kind == "report.xlsx":
        return _xlsx(start, end)
    return redirect("reports")


def _xlsx(start, end):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="C2410C")

    def sheet(ws, header, rows, money_cols=()):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
        for r in rows:
            ws.append(list(r))
        for idx in money_cols:
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2):
                for c in cell:
                    c.number_format = '"₹"#,##,##0.00'
        for i, h in enumerate(header, 1):
            ws.column_dimensions[get_column_letter(i)].width = max(10, min(40, len(str(h)) + 6))
        ws.freeze_panes = "A2"

    p = services.period(start, end)
    ws = wb.active
    ws.title = "Summary"
    ws.append([f"Lavish Shawarma — {start:%d %b %Y} to {end:%d %b %Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in [("Sales", p["sales"]), ("Orders", p["orders"]), ("Average bill", p["avg_bill"]),
                         ("Discounts", p["discounts"]), ("Expenses", p["expenses"]), ("Sales − expenses", p["profit"]),
                         ("Cancelled orders", p["cancelled"])]:
        ws.append([label, float(value) if isinstance(value, Decimal) else value])
    ws.append([])
    ws.append(["Payment method", "Amount"])
    for m, v in p["by_method"].items():
        ws.append([PaymentMethod(m).label, float(v)])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 16

    sheet(wb.create_sheet("Daily"), ["Date", "Sales", "Orders"],
          ([d.isoformat(), float(v["sales"]), v["orders"]] for d, v in p["daily"].items()), money_cols=(2,))
    sheet(wb.create_sheet("Orders"), ORDER_HEADER, _orders_rows(start, end), money_cols=(8, 9, 10, 11))
    sheet(wb.create_sheet("Items sold"), ITEM_HEADER, _item_rows(start, end), money_cols=(9, 10, 11))
    sheet(wb.create_sheet("Top items"), ["Item", "Size", "Qty", "Amount"],
          ([r["item_name"], r["variant_name"], r["qty"], float(r["amount"] or 0)] for r in p["top_items"]), money_cols=(4,))
    pays = Payment.objects.filter(business_date__range=(start, end)).select_related("order", "received_by")
    sheet(wb.create_sheet("Payments"), ["Date", "Time", "Bill", "Method", "Amount", "Reference", "Received by"],
          ([x.business_date.isoformat(), timezone.localtime(x.created_at).strftime("%H:%M"), x.order.bill_number,
            x.get_method_display(), float(x.amount), x.reference, x.received_by.username] for x in pays), money_cols=(5,))
    exps = Expense.objects.filter(business_date__range=(start, end)).select_related("created_by")
    sheet(wb.create_sheet("Expenses"), ["Date", "Category", "Description", "Amount", "Paid from", "By"],
          ([e.business_date.isoformat(), e.get_category_display(), e.description, float(e.amount),
            e.get_paid_from_display(), e.created_by.username] for e in exps), money_cols=(4,))
    counts = StockCount.objects.filter(business_date__range=(start, end)).select_related("stock_item")
    sheet(wb.create_sheet("Stock counts"), ["Date", "Item", "Unit", "Expected", "Counted", "Difference", "Value"],
          ([c.business_date.isoformat(), c.stock_item.name, c.stock_item.unit, float(c.expected_qty), float(c.counted_qty),
            float(c.variance), float(c.variance_value)] for c in counts), money_cols=(7,))
    closes = DayClose.objects.filter(business_date__range=(start, end)).order_by("business_date")
    sheet(wb.create_sheet("Day closings"),
          ["Date", "Sales", "Cash", "UPI", "Card", "Online", "Expenses", "Expected cash", "Counted cash", "Cash diff",
           "Stock loss", "Closed by"],
          ([c.business_date.isoformat(), float(c.total_sales), float(c.cash_sales), float(c.upi_sales), float(c.card_sales),
            float(c.online_sales), float(c.total_expenses), float(c.expected_cash), float(c.counted_cash),
            float(c.cash_difference), float(c.stock_loss_value), str(c.closed_by)] for c in closes),
          money_cols=tuple(range(2, 12)))

    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="lavish_report_{start}_{end}.xlsx"'
    return resp
